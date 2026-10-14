"""Render sheets of the Excel book as one HTML page (for demo screenshots).

  python scripts/book_to_html.py book.xlsx out.html [Sheet1,Sheet2] [max_rows]
"""
import html
import sys
from openpyxl import load_workbook

wb = load_workbook(sys.argv[1])
only = sys.argv[3].split(",") if len(sys.argv) > 3 else None
max_rows = int(sys.argv[4]) if len(sys.argv) > 4 else None
parts = ["<meta charset='utf-8'><style>body{font:14px Segoe UI,Arial;margin:20px;background:#fff}"
         "h2{font-size:16px;margin:18px 0 6px}table{border-collapse:collapse;max-width:1600px}"
         "td,th{border:1px solid #ccc;padding:4px 8px;vertical-align:top}"
         "th{background:#eef1f5;text-align:left}tr.w td{background:#FFE699}</style>"]
for ws in (wb[n] for n in only) if only else wb:
    parts.append(f"<h2>{html.escape(ws.title)}</h2><table>")
    rows = list(ws.iter_rows())
    for i, row in enumerate(rows[: max_rows + 1] if max_rows else rows):
        warn = any(c.fill and c.fill.fgColor and c.fill.fgColor.rgb in ("00FFE699", "FFFFE699") for c in row)
        tag = "th" if i == 0 and ws.title not in ("Estimate", "Spending") else "td"
        cells = "".join(f"<{tag}>{html.escape('' if c.value is None else str(c.value))}</{tag}>" for c in row)
        parts.append(f"<tr class='{'w' if warn else ''}'>{cells}</tr>")
    if max_rows and len(rows) > max_rows + 1:
        parts.append(f"<tr><td colspan='13'>… {len(rows) - max_rows - 1} more rows</td></tr>")
    parts.append("</table>")
open(sys.argv[2], "w", encoding="utf-8").write("".join(parts))
