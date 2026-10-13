import sys
from openpyxl import load_workbook
wb = load_workbook(sys.argv[1])
for ws in wb:
    print("==", ws.title)
    for row in ws.iter_rows(values_only=True):
        print(" | ".join("" if v is None else str(v) for v in row))
